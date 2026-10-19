"""End-to-end check of /api/finance against a fresh database (temporary, removed before merge)."""
import datetime
import json
import os
import urllib.error
import urllib.request

API = os.environ.get("API", "http://localhost:8080")
KEY = os.environ.get("SEPAY_KEY", "testkey")
VN = datetime.timezone(datetime.timedelta(hours=7))
NOW = datetime.datetime.now(VN)
MONTH = NOW.strftime("%Y-%m")


def call(method, path, body=None, headers=None, expect=200):
    data = None if body is None else json.dumps(body).encode()
    req = urllib.request.Request(API + path, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    try:
        with urllib.request.urlopen(req) as r:
            code, raw = r.status, r.read().decode()
    except urllib.error.HTTPError as e:
        code, raw = e.code, e.read().decode()
    if expect is not None and code != expect:
        raise AssertionError(f"{method} {path} -> {code} {raw[:500]}")
    return json.loads(raw) if raw else None


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print("ok -", msg)


def ts(minutes_ago):
    return (NOW - datetime.timedelta(minutes=minutes_ago)).strftime("%Y-%m-%d %H:%M:%S")


def hook(pid, typ, amount, content, accumulated, minutes_ago, key=KEY, expect=200, ref=None):
    body = {"id": pid, "gateway": "TPBank", "transactionDate": ts(minutes_ago), "accountNumber": "00001234567",
            "subAccount": "", "code": None, "content": content, "transferType": typ, "description": "",
            "transferAmount": amount, "accumulated": accumulated, "referenceCode": ref or f"FT{pid}"}
    headers = {"Authorization": f"Apikey {key}"} if key else {}
    return call("POST", "/api/finance/hooks/sepay", body, headers, expect)


def txns():
    return call("GET", f"/api/finance/transactions?month={MONTH}&size=200")["content"]


def find(desc):
    """The transaction itself, not the mirror written on the other account of a transfer."""
    hits = [t for t in txns() if (t.get("description") or "") == desc and t["source"] != "SYSTEM"]
    check(len(hits) == 1, f"exactly one transaction '{desc}' (found {len(hits)})")
    return hits[0]


def amt(x):
    return round(float(x))


def line(report, name):
    for l in report["expenses"]:
        if l["name"] == name:
            return amt(l["amount"])
    return None


# ── seeded defaults ──
accounts = {a["name"]: a for a in call("GET", "/api/finance/accounts")}
check(set(accounts) >= {"TPBank", "Techcombank", "MoMo", "Bank of America", "Tiền mặt"}, "five accounts seeded")
check(accounts["Tiền mặt"]["mode"] == "POOL" and accounts["Bank of America"]["currency"] == "USD", "cash is a pool, BoA is USD")
cats = {c["code"]: c for c in call("GET", "/api/finance/categories")}
check({"ca-phe", "tien-mat", "chi-chua-ro", "thu-chua-ro", "gia-dinh", "khac"} <= set(cats), "categories seeded")
check(len(call("GET", "/api/finance/rules")) >= 15, "rules seeded")
check(call("GET", "/api/finance/hooks/sepay")["keyRequired"] is True, "hook ping says a key is required")

# ── SePay webhook ──
r = hook(1000, "out", 1, "x", 0, 40, key="wrong", expect=401)
check(r["success"] is False, "wrong Apikey refused")
check(hook(1001, "out", 35000, "HIGHLANDS COFFEE THANH TOAN QR", 1965000, 30) == {"success": True}, "SePay answer is {success:true}")
check(hook(1001, "out", 35000, "HIGHLANDS COFFEE THANH TOAN QR", 1965000, 30) == {"success": True}, "duplicate still answered success")
hook(1002, "in", 20000000, "CONG TY NCS TRA LUONG THANG 10", 21965000, 28)
hook(1003, "out", 2000000, "RUT TIEN ATM 123", 19965000, 26)
hook(1004, "out", 500000, "NAP TIEN VI MOMO 0912", 19465000, 24)
hook(1005, "out", 120000, "NGUYEN VAN A chuyen tien", 19345000, 22)
hook(1007, "out", 50000, "GRAB 123", 19200000, 20)

hl = find("HIGHLANDS COFFEE THANH TOAN QR")
check(hl["categoryName"] == "Cà phê" and hl["status"] == "OK" and hl["accountName"] == "TPBank", "coffee by rule, on TPBank")
check(amt(hl["amount"]) == -35000 and amt(hl["balanceAfter"]) == 1965000, "amount signed, bank balance kept")
check(find("CONG TY NCS TRA LUONG THANG 10")["categoryName"] == "Lương", "salary by rule")
atm = find("RUT TIEN ATM 123")
check(atm["kind"] == "TRANSFER" and atm["transferAccountName"] == "Tiền mặt", "ATM withdrawal moves money to cash")
mirror = [t for t in txns() if t["source"] == "SYSTEM" and t["description"] == "RUT TIEN ATM 123"]
check(len(mirror) == 1 and mirror[0]["accountName"] == "Tiền mặt" and amt(mirror[0]["amount"]) == 2000000, "cash side written as a mirror")
momo_top = find("NAP TIEN VI MOMO 0912")
check(momo_top["kind"] == "TRANSFER" and momo_top["transferAccountName"] == "MoMo", "top-up moves money to MoMo")
nva = find("NGUYEN VAN A chuyen tien")
check(nva["status"] == "REVIEW" and nva["kind"] == "EXPENSE" and nva["categoryId"] is None, "unknown transfer waits for review")
check(nva["suggestedKeywords"] == "nguyen van a", f"keyword suggestion ({nva['suggestedKeywords']})")
check(find("GRAB 123")["categoryName"] == "Đi lại", "grab by rule")

accounts = {a["name"]: a for a in call("GET", "/api/finance/accounts")}
tp = accounts["TPBank"]
check(tp["accountNumber"] == "00001234567" and tp["mode"] == "FEED" and amt(tp["balance"]) == 19200000, "TPBank claimed the number, balance from the bank")
inbound = call("GET", "/api/finance/inbound?size=50")["content"]
check(any(i["status"] == "DUPLICATE" for i in inbound) and sum(i["status"] == "SAVED" for i in inbound) == 6, "webhook log: 6 saved, 1 duplicate")

# ── quick notes ──
cash_id, momo_id, tcb_id = accounts["Tiền mặt"]["id"], accounts["MoMo"]["id"], accounts["Techcombank"]["id"]
n1 = call("POST", "/api/finance/notes", {"text": "cafe", "amount": -35000})
check(n1["accountName"] == "Tiền mặt" and n1["categoryName"] == "Cà phê" and n1["status"] == "OK", "note defaults to cash, rule applies")
n2 = call("POST", "/api/finance/notes", {"text": "bún chả", "amount": -45000, "accountId": momo_id})
check(n2["categoryName"] == "Ăn uống" and n2["accountName"] == "MoMo", "accented note matched (bún chả → Ăn uống)")
n3 = call("POST", "/api/finance/notes", {"text": "linh tinh", "amount": -20000})
check(n3["status"] == "REVIEW", "unknown note waits for review")
n4 = call("POST", "/api/finance/notes", {"text": "cafe", "amount": -35000, "occurredAt": NOW.strftime("%Y-%m-%dT%H:%M:%S+07:00"), "categoryId": cats["ca-phe"]["id"]})
check(n4["categorizedBy"] == "USER" and n4["status"] == "OK", "note with a chosen category and an offset time")
call("DELETE", f"/api/finance/transactions/{n4['id']}", expect=204)
chips = call("GET", "/api/finance/notes/chips")
check(any(c["text"] == "cafe" for c in chips), "quick chips built from notes")

# ── balance check-ins (MoMo) ──
m = call("POST", f"/api/finance/accounts/{momo_id}/balance", {"balance": 300000})
check(amt(m["balance"]) == 300000, "first check-in sets the start")
m = call("POST", f"/api/finance/accounts/{momo_id}/balance", {"balance": 200000})
adj = find("Chênh lệch khi chốt số dư")
check(amt(adj["amount"]) == -100000 and adj["categoryName"] == "Chi chưa rõ", "second check-in: gap becomes 'Chi chưa rõ'")
call("POST", f"/api/finance/accounts/{cash_id}/balance", {"balance": 1}, expect=400)

# ── the month report ──
rep = call("GET", f"/api/finance/report?month={MONTH}")
check(amt(rep["expense"]) == 2350000, f"spending counts cash at withdrawal (got {rep['expense']})")
check(amt(rep["income"]) == 20000000, "income")
check(line(rep, "Cà phê") == 70000 and line(rep, "Ăn uống") == 45000 and line(rep, "Đi lại") == 50000, "category lines")
check(line(rep, "Chưa phân loại") == 140000 and line(rep, "Chi chưa rõ") == 100000, "uncategorized and unknown lines")
check(line(rep, "Tiền mặt lặt vặt") == 1945000, f"cash bucket = withdrawn - noted (got {line(rep, 'Tiền mặt lặt vặt')})")
check(sum(amt(l["amount"]) for l in rep["expenses"]) == amt(rep["expense"]), "lines add up to the total")
check(len(rep["gaps"]) == 1 and amt(rep["gaps"][0]["missing"]) == -95000, "balance gap found (-95k)")

call("PUT", "/api/finance/settings", {"flexibleBudget": 5000000})
rep = call("GET", f"/api/finance/report?month={MONTH}")
b = rep["budget"]
check(amt(b["limit"]) == 5000000 and amt(b["spent"]) == 2350000 and amt(b["remaining"]) == 2650000 and b["daysLeft"] >= 1, "budget")

# ── learning from a correction ──
u = call("PUT", f"/api/finance/transactions/{nva['id']}", {"categoryId": cats["gia-dinh"]["id"], "learnKeywords": "nguyen van a"})
check(u["categoryName"] == "Gia đình" and u["status"] == "OK" and u["categorizedBy"] == "USER", "corrected by hand")
check(any(r["keywords"] == "nguyen van a" and r["origin"] == "USER" for r in call("GET", "/api/finance/rules")), "rule learned")
hook(1008, "out", 100000, "NGUYEN VAN A chuyen tien lan 2", 19100000, 18)
check(find("NGUYEN VAN A chuyen tien lan 2")["categoryName"] == "Gia đình", "learned rule used next time")

# ── deleting a transfer takes its mirror with it ──
call("DELETE", f"/api/finance/transactions/{atm['id']}", expect=204)
check(not any((t.get("description") or "") == "RUT TIEN ATM 123" for t in txns()), "withdrawal and its cash side deleted")
rep = call("GET", f"/api/finance/report?month={MONTH}")
check(amt(rep["expense"]) == 450000, f"report after delete (got {rep['expense']})")
check(sum(amt(l["amount"]) for l in rep["expenses"]) == 450000, "lines still add up")

# ── manual transaction, transfer and back ──
t = call("POST", "/api/finance/transactions", {"accountId": tcb_id, "amount": -99000, "description": "netflix"})
check(t["categoryName"] == "Đăng ký dịch vụ", "manual entry goes through rules")
t = call("PUT", f"/api/finance/transactions/{t['id']}", {"kind": "TRANSFER", "transferAccountId": momo_id})
check(t["kind"] == "TRANSFER" and t["transferAccountName"] == "MoMo", "turned into a transfer")
check(amt({a["name"]: a for a in call("GET", "/api/finance/accounts")}["MoMo"]["balance"]) == 299000, "MoMo running balance moved with the mirror")
t = call("PUT", f"/api/finance/transactions/{t['id']}", {"categoryId": cats["dang-ky"]["id"]})
check(t["kind"] == "EXPENSE" and t["transferPeerId"] is None, "back to spending")
check(amt({a["name"]: a for a in call("GET", "/api/finance/accounts")}["MoMo"]["balance"]) == 200000, "mirror removed, balance restored")

# ── categories, rules, accounts ──
c = call("POST", "/api/finance/categories", {"name": "Thú cưng", "icon": "🐶"})
check(c["code"] == "thu-cung" and c["kind"] == "EXPENSE" and c["budgetGroup"] == "FLEXIBLE", "category created")
call("DELETE", f"/api/finance/categories/{c['id']}", expect=204)
call("DELETE", f"/api/finance/categories/{cats['tien-mat']['id']}", expect=400)
rule = call("POST", "/api/finance/rules", {"keywords": "linh tinh", "categoryId": cats["khac"]["id"]})
check(call("POST", "/api/finance/rules/apply")["updated"] == 1, "rules applied to the waiting note")
check(find("linh tinh")["categoryName"] == "Khác", "waiting note categorized")
call("DELETE", f"/api/finance/rules/{rule['id']}", expect=204)
a = call("POST", "/api/finance/accounts", {"name": "Ví phụ", "mode": "RECONCILE"})
call("DELETE", f"/api/finance/accounts/{a['id']}", expect=204)
call("DELETE", f"/api/finance/accounts/{tp['id']}", expect=400)
call("POST", "/api/finance/transactions/confirm-all")
check(len(call("GET", "/api/finance/review")) == 0, "review list empty after confirm-all")

# ── summary and trend ──
s = call("GET", "/api/finance/summary")
check(amt(s["netWorthVnd"]) == 19100000 + 200000, f"net worth = TPBank + MoMo (got {s['netWorthVnd']})")
tr = call("GET", "/api/finance/trend?months=6")
check(tr[-1]["month"] == MONTH and amt(tr[-1]["expense"]) == amt(call("GET", f"/api/finance/report?month={MONTH}")["expense"]), "trend ends with this month")
check(amt(tr[-1]["netWorth"]) == 19300000, "trend net worth now")
st = call("GET", "/api/finance/settings")
check(st["sepayKeySet"] is True and amt(st["flexibleBudget"]) == 5000000, "settings")
call("PUT", "/api/finance/settings", {"usdRate": 26000})
boa = {a["name"]: a for a in call("GET", "/api/finance/accounts")}["Bank of America"]["id"]
call("POST", f"/api/finance/accounts/{boa}/balance", {"balance": 100})
check(amt(call("GET", "/api/finance/summary")["netWorthVnd"]) == 19300000 + 2600000, "USD balance converted at the manual rate")
# ── a learned phrase is the words as written, so the next note with them matches ──
n = call("POST", "/api/finance/notes", {"text": "đồ chơi cho bé", "amount": -420000})
check(n["status"] == "REVIEW" and n["suggestedKeywords"] == "do choi", f"phrase kept contiguous (got {n['suggestedKeywords']!r})")
call("PUT", f"/api/finance/transactions/{n['id']}", {"categoryId": cats["mua-sam"]["id"], "learnKeywords": n["suggestedKeywords"]})
n = call("POST", "/api/finance/notes", {"text": "đồ chơi lego", "amount": -300000})
check(n["categoryName"] == "Mua sắm" and n["status"] == "OK", "learned phrase matches the next note")

# ── stamps are UTC whatever the server's zone (this run starts the API with TZ=Asia/Ho_Chi_Minh) ──
last = call("GET", "/api/finance/inbound?size=1")["content"][0]["createdAt"]
utc_now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
off = abs((utc_now - datetime.datetime.fromisoformat(last[:19])).total_seconds())
check(off < 3600, f"webhook log stamped in UTC (off by {off:.0f}s)")
# ── bank statement import, counterparties ──
import io, uuid
from openpyxl import Workbook


def statement(rows, account="00001234567"):
    """A TPBank-style 'Lịch sử giao dịch tài khoản' export: title rows, then the header on row 9, newest first."""
    wb = Workbook(); ws = wb.active; ws.title = "VN"
    ws["A3"] = "LỊCH SỬ GIAO DỊCH TÀI KHOẢN"
    ws["A5"], ws["B5"] = "Thời gian:", "01/10/2026 - 19/10/2026"
    ws["A6"], ws["B6"] = "Tài khoản giao dịch:", account
    ws.append([]); ws.append([])
    ws.append(["Ngày thực hiện", "Ngày hiệu lực", "Mô tả giao dịch", "Ghi nợ", "Ghi có", "Số dư", "Tài khoản đối ứng", "Tên tài khoản", "Mã giao dịch"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def upload(path, data, expect=200):
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"stmt.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() + data + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(API + path, data=body, method="POST")
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    try:
        with urllib.request.urlopen(req) as r:
            code, raw = r.status, r.read().decode()
    except urllib.error.HTTPError as e:
        code, raw = e.code, e.read().decode()
    if code != expect:
        raise AssertionError(f"POST {path} -> {code} {raw[:500]}")
    return json.loads(raw)


def vn(days=0, minutes=0):
    return (NOW - datetime.timedelta(days=days, minutes=minutes)).strftime("%d-%m-%Y %H:%M:%S")


def by_desc(desc, account=None):
    return [t for t in call("GET", "/api/finance/transactions?size=200&q=" + urllib.parse.quote(desc))["content"]
            if (t.get("description") == desc or t.get("bankDescription") == desc) and t["source"] != "SYSTEM"
            and (account is None or t["accountName"] == account)]


import urllib.parse
accounts = {a["name"]: a for a in call("GET", "/api/finance/accounts")}
call("PUT", f"/api/finance/accounts/{accounts['Techcombank']['id']}", {"accountNumber": "1903000111"})
momo_before = amt(accounts["MoMo"]["balance"])
hook(2001, "out", 85000, "Vu Ngoc Anh chuyen tien QR", 0, 12)
hook(2002, "out", 21000, "QRBOXVVB077393", 0, 11)
G = "GIAO DICH THE 0465 TAI Google YouTubePremium650-2530000US"
rows = [
    [vn(minutes=11), vn(), "QRBOXVVB077393", 21000, "", 1, "01431140003", "CT TNHH TMAI VA DICH VU ATT VN", "FT2002"],
    [vn(minutes=12), vn(), "Vu Ngoc Anh chuyen tien QR", 85000, "", 1, "0987654321", "TRAN THI B", "FT2001"],
    [vn(days=1), vn(), "TRAN THI B tra tien", "", 30000, 1, "0987654321", "TRAN THI B", "011CT00000000009"],
    [vn(days=2), vn(), "Vu Ngoc Anh chuyen tien QR", 300000, "", 1, "5550001", "LE VAN C", "011V600000000008"],
    [vn(days=3), vn(), "Vu Ngoc Anh chuyen tien QR", 50000, "", 1, "0987654321", "TRAN THI B", "011V600000000001"],
    [vn(days=4), vn(), "GIAO DICH THE 0465 TAI Green SM8424393666VN", 61000, "", 1, "", "", "000ZP0000000001"],
    [vn(days=4, minutes=5), vn(), G, 11000, "", 1, "", "", "000ZP0000000002"],
    [vn(days=4, minutes=5), vn(), G, 165556, "", 1, "", "", "000ZP0000000002"],
    [vn(days=5), vn(), "VU NGOC ANH chuyen tien", 10000, "", 1, "1903000111", "VU NGOC ANH", "011CTMB00000003"],
    [vn(days=5, minutes=30), vn(), "NAP TIEN VI MOMO - 0912345678 - 123456", 200000, "", 1, "99MM2610000001", "CTY CP DICH VU DI DONG TRUC TUYEN", "011V600000000004"],
    [(NOW - datetime.timedelta(days=6)).strftime("%d/%m/%Y"), vn(), "TRA LAI TIEN GUI TK: 00001234567", "", 4471, 1, "", "", "011CT00000000005"],
]
file = statement(rows)
count_before = call("GET", "/api/finance/transactions?size=1")["totalElements"]
p = upload("/api/finance/import/preview", file)
check(p["rows"] == 11 and p["matched"] == 2 and p["enriched"] == 2 and p["created"] == 9 and p["already"] == 0 and not p["committed"],
      f"preview: 11 rows, 2 from the feed, 9 new ({p})")
check(p["accountName"] == "TPBank" and p["withCounterparty"] == 10, "preview finds the account by number; all but the interest line name a counterparty")
check(call("GET", "/api/finance/transactions?size=1")["totalElements"] == count_before, "preview writes nothing")
c = upload("/api/finance/import", file)
check(c["committed"] and c["created"] == 9 and c["matched"] == 2, "import: 9 new lines, 2 matched")
again = upload("/api/finance/import", file)
check(again["created"] == 0 and again["already"] == 9 and again["matched"] == 2, "importing the same file again adds nothing")

f1 = by_desc("Vu Ngoc Anh chuyen tien QR")
check(len(f1) == 3 and all(t["counterparty"] in ("TRAN THI B", "LE VAN C") for t in f1), "generic memos got their counterparties")
sp = [t for t in f1 if t["source"] == "SEPAY"][0]
check(sp["counterparty"] == "TRAN THI B" and sp["counterpartyAccount"] == "0987654321", "the feed line got the counterparty from the statement")
gr = by_desc("GIAO DICH THE 0465 TAI Green SM8424393666VN")[0]
check(gr["counterparty"] == "Green SM" and gr["categoryName"] == "Đi lại" and gr["source"] == "STATEMENT", "card merchant read from the memo, upgraded rule placed it")
yt = by_desc(G)
check(len(yt) == 2 and {amt(t["amount"]) for t in yt} == {-11000, -165556} and all(t["categoryName"] == "Đăng ký dịch vụ" for t in yt),
      "card payment and its fee share a code: both kept")
own = by_desc("VU NGOC ANH chuyen tien")[0]
check(own["kind"] == "TRANSFER" and own["transferAccountName"] == "Techcombank", "money to my own account number is a move, not spending")
mt = by_desc("NAP TIEN VI MOMO - 0912345678 - 123456")[0]
check(mt["kind"] == "TRANSFER" and mt["transferAccountName"] == "MoMo", "old MoMo top-up is a move to MoMo")
check(amt({a["name"]: a for a in call("GET", "/api/finance/accounts")}["MoMo"]["balance"]) == momo_before,
      "a line older than the last check-in leaves the checked balance alone")
it = by_desc("TRA LAI TIEN GUI TK: 00001234567")[0]
check(it["categoryName"] == "Lãi & hoàn tiền" and amt(it["amount"]) == 4471, "date-only interest line imported as income")

groups = {g["name"]: g for g in call("GET", "/api/finance/counterparties")}
b = groups.get("TRAN THI B")
check(b is not None and b["waiting"] == 3 and b["outCount"] == 2 and b["inCount"] == 1, f"TRAN THI B: 3 lines waiting, 2 out 1 in ({b})")
check("CT TNHH TMAI VA DICH VU ATT VN" in groups and "Green SM" not in groups, "groups list only who still waits")
res = call("POST", "/api/finance/counterparties/assign", {"name": "TRAN THI B", "categoryId": cats["gia-dinh"]["id"], "label": "Gửi B"})
check(res["updated"] == 3, f"one decision placed all 3 lines ({res})")
lb = by_desc("Gửi B")
check(len(lb) == 3 and all(t["categoryName"] == "Gia đình" and t["status"] == "OK" for t in lb), "they are Gia đình and labelled")
check(sorted(t["bankDescription"] for t in lb) == ["TRAN THI B tra tien", "Vu Ngoc Anh chuyen tien QR", "Vu Ngoc Anh chuyen tien QR"], "the memo is kept as the bank line")
rule = [r for r in call("GET", "/api/finance/rules") if r["counterpartyName"] == "TRAN THI B"]
check(len(rule) == 1 and rule[0]["label"] == "Gửi B" and rule[0]["direction"] == "ANY" and rule[0]["priority"] == 5, "a counterparty rule was saved")
res = call("POST", "/api/finance/counterparties/assign", {"name": "CT TNHH TMAI VA DICH VU ATT VN", "categoryId": cats["an-uong"]["id"]})
check(res["updated"] == 1, "ATT placed")

hook(2004, "out", 35000, "QRBOXVVB077393", 0, 5)
q = [t for t in by_desc("QRBOXVVB077393") if t["source"] == "SEPAY" and amt(t["amount"]) == -35000][0]
check(q["counterparty"] == "CT TNHH TMAI VA DICH VU ATT VN" and q["categoryName"] == "Ăn uống", "a shop's fixed QR is known live: counterparty and category at once")
hook(2005, "out", 40000, "Vu Ngoc Anh chuyen tien QR", 0, 4)
g5 = [t for t in call("GET", "/api/finance/transactions?size=200")["content"] if t["description"] == "Vu Ngoc Anh chuyen tien QR" and t["source"] == "SEPAY"]
check(len(g5) == 1 and g5[0]["counterparty"] is None and g5[0]["status"] == "REVIEW", "a generic memo does not guess a counterparty")
hook(2006, "out", 50000, "Vu Ngoc Anh chuyen tien QR", 0, 3, ref="011V600000000001")
check(call("GET", "/api/finance/inbound?size=1")["content"][0]["status"] == "DUPLICATE", "a feed line already imported from a statement is not added twice")

c_line = [t for t in by_desc("Vu Ngoc Anh chuyen tien QR") if t["counterparty"] == "LE VAN C"][0]
u = call("PUT", f"/api/finance/transactions/{c_line['id']}", {"categoryId": cats["qua-hieu-hi"]["id"], "learnCounterparty": True, "label": "Quà C"})
check(u["description"] == "Quà C" and u["categoryName"] == "Quà & hiếu hỉ", "edit sheet: category and label")
check(any(r["counterpartyName"] == "LE VAN C" and r["label"] == "Quà C" for r in call("GET", "/api/finance/rules")), "edit sheet: remembered for everything with LE VAN C")
bad = statement([[vn(), vn(), "x", 1000, "", 1, "", "", "R1"]], account="999999999")
err = upload("/api/finance/import/preview", bad, expect=400)
check("••9999" in err["error"], f"unknown account number explained ({err['error']})")

print("\nALL FINANCE CHECKS PASSED")
